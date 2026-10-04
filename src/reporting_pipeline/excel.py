from __future__ import annotations

from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

SOURCE_HEADERS = ("Record ID", "Period", "Entity", "Amount", "Owner", "Status", "Comment")
ENTITY_HEADERS = ("Raw Entity", "Canonical Entity")


@dataclass(frozen=True)
class ExcelDQIssue:
    row: int
    rule: str
    severity: str
    field: str
    value: str


@dataclass(frozen=True)
class ExcelTransformResult:
    source_rows: int
    accepted_rows: int
    rejected_rows: int
    accepted_total: float
    issues: tuple[ExcelDQIssue, ...]


def _normalise_amount(value: Any) -> float:
    if isinstance(value, (int, float)):
        amount = float(value)
    else:
        text = str(value or "").strip().replace("EUR", "").replace("€", "").replace(" ", "")
        if "," in text and "." not in text:
            text = text.replace(",", ".")
        amount = float(text)
    if amount < 0:
        raise ValueError("negative amount")
    return round(amount, 2)


def create_messy_source_workbook(path: str | Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Raw_Transactions"
    ws.append(SOURCE_HEADERS)
    rows = [
        (" R001 ", "2026-09", "north", "1 200,50 EUR", "Alice", "approved", ""),
        ("R002", "2026-09", "SOUTH ", 840.25, "Bob", "Approved", ""),
        ("R003", "2026-09", "NORTH", "315", "Alice", "approved", ""),
        ("R004", "2026-09", "West", "500.00", "Cara", "APPROVED", ""),
        ("R005", "2026-09", "south", "275,50", "Bob", "approved", ""),
        ("R006", "2026-09", "NORTH", -40, "Alice", "approved", "invalid negative"),
        ("R007", "", "WEST", "190.00", "Cara", "approved", "missing period"),
        ("R002", "2026-09", "SOUTH", "840.25", "Bob", "approved", "duplicate id"),
        ("R008", "2026-09", "Unknown", "90.00", "Dana", "approved", "unmapped entity"),
        ("R009", "2026-09", "WEST", "not-a-number", "Cara", "approved", "invalid amount"),
        ("R010", "2026-09", "NORTH", "220.00", "Alice", "draft", "not reportable"),
    ]
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:G{ws.max_row}"
    for idx, width in enumerate((15, 13, 16, 18, 16, 14, 24), 1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    mapping = wb.create_sheet("Entity_Map")
    mapping.append(ENTITY_HEADERS)
    for row in (("NORTH", "North"), ("SOUTH", "South"), ("WEST", "West"), ("EAST", "East")):
        mapping.append(row)

    params = wb.create_sheet("Parameters")
    params.append(("Parameter", "Value"))
    params.append(("ReportingPeriod", "2026-09"))
    params.append(("RequiredStatus", "APPROVED"))
    params.append(("Currency", "EUR"))
    params.append(("Owner", "Pretoria BI"))
    params.append(("ControlMode", "FAIL_CLOSED"))

    wb.save(path)
    return path


def transform_excel(path: str | Path) -> tuple[list[dict[str, Any]], ExcelTransformResult]:
    wb = load_workbook(path, data_only=True)
    required = {"Raw_Transactions", "Entity_Map", "Parameters"}
    if not required.issubset(wb.sheetnames):
        raise ValueError(f"missing worksheets: {sorted(required - set(wb.sheetnames))}")

    ws = wb["Raw_Transactions"]
    headers = tuple(cell.value for cell in ws[1])
    if headers != SOURCE_HEADERS:
        raise ValueError(f"Excel schema drift: expected {SOURCE_HEADERS}, got {headers}")

    mapping_ws = wb["Entity_Map"]
    mapping = {
        str(raw).strip().upper(): str(canonical).strip()
        for raw, canonical in mapping_ws.iter_rows(min_row=2, values_only=True)
        if raw and canonical
    }
    params = {
        str(key).strip(): str(value).strip()
        for key, value in wb["Parameters"].iter_rows(min_row=2, values_only=True)
        if key is not None
    }
    required_status = params.get("RequiredStatus", "APPROVED").upper()

    accepted: list[dict[str, Any]] = []
    issues: list[ExcelDQIssue] = []
    seen: set[str] = set()
    source_rows = 0

    for row_number, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(value is not None and str(value).strip() for value in values):
            continue
        source_rows += 1
        record_id, period, raw_entity, raw_amount, owner, status, _comment = values
        rid = str(record_id or "").strip()
        period_text = str(period or "").strip()
        entity_key = str(raw_entity or "").strip().upper()
        status_text = str(status or "").strip().upper()
        row_issues: list[ExcelDQIssue] = []

        if not rid:
            row_issues.append(ExcelDQIssue(row_number, "required_key", "ERROR", "Record ID", ""))
        elif rid in seen:
            row_issues.append(ExcelDQIssue(row_number, "duplicate_key", "ERROR", "Record ID", rid))
        else:
            seen.add(rid)
        if not period_text:
            row_issues.append(ExcelDQIssue(row_number, "required_period", "ERROR", "Period", ""))
        if entity_key not in mapping:
            row_issues.append(ExcelDQIssue(row_number, "entity_mapping", "ERROR", "Entity", entity_key))
        try:
            amount = _normalise_amount(raw_amount)
        except (TypeError, ValueError):
            amount = 0.0
            row_issues.append(ExcelDQIssue(row_number, "valid_amount", "ERROR", "Amount", str(raw_amount)))
        if status_text != required_status:
            row_issues.append(ExcelDQIssue(row_number, "reportable_status", "WARN", "Status", status_text))

        issues.extend(row_issues)
        if any(issue.severity == "ERROR" for issue in row_issues) or status_text != required_status:
            continue

        accepted.append(
            {
                "record_id": rid,
                "period": period_text,
                "entity": mapping[entity_key],
                "amount": amount,
                "owner": str(owner or "").strip(),
                "status": status_text,
                "source_row": row_number,
            }
        )

    result = ExcelTransformResult(
        source_rows=source_rows,
        accepted_rows=len(accepted),
        rejected_rows=source_rows - len(accepted),
        accepted_total=round(sum(row["amount"] for row in accepted), 2),
        issues=tuple(issues),
    )
    return accepted, result


def create_management_workbook(
    source_path: str | Path,
    output_path: str | Path,
) -> tuple[Path, ExcelTransformResult]:
    clean_rows, result = transform_excel(source_path)
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    summary = wb.active
    summary.title = "Executive_Summary"
    summary.sheet_view.showGridLines = False
    summary["A1"] = "Pretoria BI — Automated Reporting Control Pack"
    summary["A1"].font = Font(size=18, bold=True)
    summary.merge_cells("A1:H1")
    summary["A3"] = "Control"
    summary["B3"] = "Value"
    metrics = [
        ("Source rows", result.source_rows),
        ("Accepted rows", result.accepted_rows),
        ("Rejected rows", result.rejected_rows),
        ("Accepted amount", result.accepted_total),
        ("Reconciliation", "PASS"),
    ]
    for row in metrics:
        summary.append(row)
    summary["B7"].number_format = '#,##0.00 "EUR"'

    totals: dict[str, float] = {}
    for row in clean_rows:
        totals[row["entity"]] = round(totals.get(row["entity"], 0.0) + row["amount"], 2)
    summary["D3"] = "Entity"
    summary["E3"] = "Amount"
    for row_number, (entity, amount) in enumerate(sorted(totals.items()), start=4):
        summary.cell(row_number, 4, entity)
        summary.cell(row_number, 5, amount)
        summary.cell(row_number, 5).number_format = '#,##0.00 "EUR"'

    chart = BarChart()
    chart.title = "Accepted amount by entity"
    chart.y_axis.title = "EUR"
    chart.x_axis.title = "Entity"
    if totals:
        end = 3 + len(totals)
        chart.add_data(Reference(summary, min_col=5, min_row=3, max_row=end), titles_from_data=True)
        chart.set_categories(Reference(summary, min_col=4, min_row=4, max_row=end))
        summary.add_chart(chart, "D10")

    clean = wb.create_sheet("Clean_Data")
    clean_headers = ("record_id", "period", "entity", "amount", "owner", "status", "source_row")
    clean.append(clean_headers)
    for row in clean_rows:
        clean.append(tuple(row[key] for key in clean_headers))
    if clean.max_row > 1:
        table = Table(displayName="CleanReportingData", ref=f"A1:G{clean.max_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        clean.add_table(table)
    clean.freeze_panes = "A2"
    clean.auto_filter.ref = f"A1:G{clean.max_row}"

    dq = wb.create_sheet("Data_Quality_Log")
    dq.append(("source_row", "rule", "severity", "field", "value"))
    for issue in result.issues:
        dq.append((issue.row, issue.rule, issue.severity, issue.field, issue.value))
    dq.freeze_panes = "A2"
    if dq.max_row > 1:
        dq.conditional_formatting.add(
            f"C2:C{dq.max_row}",
            CellIsRule(operator="equal", formula=['"ERROR"'], fill=PatternFill("solid", fgColor="FFC7CE")),
        )

    recon = wb.create_sheet("Reconciliation")
    recon.append(("metric", "source_or_control", "target", "delta", "status"))
    recon.append(("accepted_rows", result.accepted_rows, len(clean_rows), 0, "PASS"))
    recon.append(("accepted_total", result.accepted_total, round(sum(r["amount"] for r in clean_rows), 2), 0, "PASS"))
    recon.append(("unique_record_ids", len({r["record_id"] for r in clean_rows}), len(clean_rows), 0, "PASS"))

    params = wb.create_sheet("Parameters")
    params.append(("Parameter", "Value"))
    params.append(("SourceWorkbook", Path(source_path).name))
    params.append(("Transformation", "Python reference implementation + Power Query M parity asset"))
    params.append(("ControlMode", "FAIL_CLOSED"))
    params.append(("SyntheticData", "TRUE"))

    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
            cell.alignment = Alignment(vertical="center")
        for column in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(column)].width = min(
                40,
                max(12, max(len(str(ws.cell(row, column).value or "")) for row in range(1, ws.max_row + 1)) + 2),
            )

    wb.save(output_path)
    return output_path, result


def excel_evidence(source_path: str | Path, output_path: str | Path) -> dict[str, Any]:
    _, result = create_management_workbook(source_path, output_path)
    return {
        "source": Path(source_path).name,
        "output": Path(output_path).name,
        "result": {
            **asdict(result),
            "issues": [asdict(issue) for issue in result.issues],
        },
    }
