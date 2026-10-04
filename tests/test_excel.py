from pathlib import Path

from openpyxl import load_workbook

from reporting_pipeline.excel import (
    SOURCE_HEADERS,
    create_management_workbook,
    create_messy_source_workbook,
    transform_excel,
)


def test_messy_workbook_contract_and_transformation(tmp_path):
    source = create_messy_source_workbook(tmp_path / "source.xlsx")
    rows, result = transform_excel(source)

    assert result.source_rows == 11
    assert result.accepted_rows == 6
    assert result.rejected_rows == 5
    assert result.accepted_total == 3351.25
    assert len({row["record_id"] for row in rows}) == len(rows)
    assert {row["entity"] for row in rows} == {"North", "South", "West"}


def test_management_workbook_is_openable_and_reconciled(tmp_path):
    source = create_messy_source_workbook(tmp_path / "source.xlsx")
    output, result = create_management_workbook(source, tmp_path / "management.xlsx")
    wb = load_workbook(output, data_only=False)

    assert wb.sheetnames == [
        "Executive_Summary",
        "Clean_Data",
        "Data_Quality_Log",
        "Reconciliation",
        "Parameters",
    ]
    assert wb["Executive_Summary"]["B7"].value == result.accepted_total
    statuses = [row[4] for row in wb["Reconciliation"].iter_rows(min_row=2, values_only=True)]
    assert statuses == ["PASS", "PASS", "PASS"]


def test_excel_schema_drift_fails_closed(tmp_path):
    source = create_messy_source_workbook(tmp_path / "source.xlsx")
    wb = load_workbook(source)
    wb["Raw_Transactions"]["A1"] = "Unexpected Key"
    wb.save(source)

    try:
        transform_excel(source)
    except ValueError as exc:
        assert "schema drift" in str(exc)
    else:
        raise AssertionError("schema drift was accepted")


def test_duplicate_and_invalid_rows_are_logged(tmp_path):
    source = create_messy_source_workbook(tmp_path / "source.xlsx")
    _, result = transform_excel(source)
    rules = {issue.rule for issue in result.issues}

    assert {"duplicate_key", "required_period", "entity_mapping", "valid_amount", "reportable_status"} <= rules
    assert any(issue.rule == "valid_amount" and issue.severity == "ERROR" for issue in result.issues)


def test_source_headers_are_explicit():
    assert SOURCE_HEADERS == (
        "Record ID",
        "Period",
        "Entity",
        "Amount",
        "Owner",
        "Status",
        "Comment",
    )
